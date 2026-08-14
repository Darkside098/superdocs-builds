from pathlib import Path

from docx import Document as DocxDocument

from superdocs_template_inference.ingestion import DOCXLoader
from superdocs_template_inference.models import Document, ParagraphBlock
from superdocs_template_inference.profiling import DocumentProfiler
from superdocs_template_inference.template_inference import TemplateInferer
from superdocs_template_inference.template_inference.result import (
    ConditionalRule,
    TemplateInferenceResult,
    TemplateSection,
    VariableField,
)
from superdocs_template_inference.template_inference.regeneration import DocxRegenerator


def _build_template() -> TemplateInferenceResult:
    return TemplateInferenceResult(
        family_id="synthetic_family",
        document_ids=["doc-1"],
        confidence=1.0,
        sections=[
            TemplateSection(
                section_id="section_0",
                title_or_pattern="company_header",
                inferred_type="REQUIRED",
                order=0,
                presence_frequency=1.0,
                presence_type="always",
                content_representation="Welcome to {{company_name}}.",
            ),
            TemplateSection(
                section_id="section_1",
                title_or_pattern="remote_work_setup",
                inferred_type="CONDITIONAL",
                order=1,
                presence_frequency=0.5,
                presence_type="conditional",
                content_representation="Remote work setup: {{employment_type}}.",
                conditional_info={"condition": {"field": "employment_type", "operator": "==", "value": "Remote"}},
            ),
            TemplateSection(
                section_id="section_2",
                title_or_pattern="closing",
                inferred_type="REQUIRED",
                order=2,
                presence_frequency=1.0,
                presence_type="always",
                content_representation="Sincerely, {{manager_name}}.",
            ),
        ],
        variables=[
            VariableField(
                variable_name="company_name",
                semantic_role="company",
                section_context="company_header",
                observed_values=["Contoso Labs"],
                frequency=1.0,
                confidence=1.0,
            ),
            VariableField(
                variable_name="employment_type",
                semantic_role="employment_type",
                section_context="remote_work_setup",
                observed_values=["Remote", "Hybrid"],
                frequency=1.0,
                confidence=1.0,
            ),
            VariableField(
                variable_name="manager_name",
                semantic_role="person_name",
                section_context="closing",
                observed_values=["Alice Johnson"],
                frequency=1.0,
                confidence=1.0,
            ),
        ],
        conditional_rules=[
            ConditionalRule(
                target_section="remote_work_setup",
                condition={"field": "employment_type", "operator": "==", "value": "Remote"},
                confidence=0.95,
                status="inferred",
            )
        ],
    )


def test_regeneration_substitutes_variables_and_keeps_order(tmp_path):
    template = _build_template()
    values = {"company_name": "Contoso Labs", "employment_type": "Remote", "manager_name": "Alice Johnson"}
    output_path = tmp_path / "synthetic_regenerated.docx"

    result = DocxRegenerator().regenerate(template, values, output_path=str(output_path))

    assert result.succeeded is True
    assert result.errors == []
    assert result.included_sections == ["company_header", "remote_work_setup", "closing"]

    doc = DocxDocument(str(output_path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Welcome to Contoso Labs." in text
    assert "Remote work setup: Remote." in text
    assert "Sincerely, Alice Johnson." in text
    assert text.index("Welcome to Contoso Labs.") < text.index("Remote work setup: Remote.") < text.index("Sincerely, Alice Johnson.")


def test_regeneration_omits_conditional_section_when_rule_is_false(tmp_path):
    template = _build_template()
    values = {"company_name": "Contoso Labs", "employment_type": "Hybrid", "manager_name": "Alice Johnson"}
    output_path = tmp_path / "conditional_false.docx"

    result = DocxRegenerator().regenerate(template, values, output_path=str(output_path))

    assert result.succeeded is True
    assert "remote_work_setup" not in result.included_sections
    doc = DocxDocument(str(output_path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Remote work setup" not in text
    assert "Welcome to Contoso Labs." in text
    assert "Sincerely, Alice Johnson." in text


def test_regeneration_returns_explicit_error_for_missing_variable_value(tmp_path):
    template = _build_template()
    values = {"employment_type": "Remote", "manager_name": "Alice Johnson"}
    output_path = tmp_path / "missing_value.docx"

    result = DocxRegenerator().regenerate(template, values, output_path=str(output_path))

    assert result.succeeded is False
    assert any("Missing variable value" in err for err in result.errors)
    assert "company_name" in result.errors[0]


def test_regeneration_warns_when_template_lacks_safe_content(tmp_path):
    template = _build_template()
    template.sections[0].content_representation = ""
    values = {"company_name": "Contoso Labs", "employment_type": "Remote", "manager_name": "Alice Johnson"}
    output_path = tmp_path / "insufficient_template.docx"

    result = DocxRegenerator().regenerate(template, values, output_path=str(output_path))

    assert result.succeeded is False
    assert any("insufficient information" in warning.lower() for warning in result.warnings)


def test_regeneration_uses_source_sections_when_available(tmp_path):
    source_document = Document(
        document_id="doc-1",
        filename="sample.docx",
        file_type="docx",
        loaded_at="2026-01-01T00:00:00Z",
        blocks=[
            ParagraphBlock(text="company_header", is_heading=True, heading_level=1),
            ParagraphBlock(text="Welcome to Bluebird Labs."),
            ParagraphBlock(text="closing", is_heading=True, heading_level=1),
            ParagraphBlock(text="Sincerely, Dana Lee."),
        ],
    )
    template = _build_template()
    template.sections[0].content_representation = ""
    template.sections[2].content_representation = ""
    values = {"company_name": "Bluebird Labs", "employment_type": "Remote", "manager_name": "Dana Lee"}
    output_path = tmp_path / "source_based.docx"

    result = DocxRegenerator().regenerate(template, values, source_document=source_document, output_path=str(output_path))

    assert result.succeeded is True
    doc = DocxDocument(str(output_path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Welcome to Bluebird Labs." in text
    assert "Sincerely, Dana Lee." in text


def test_integration_regeneration_uses_real_benchmark_docx(tmp_path):
    doc_path = next(Path("corpus").glob("**/*.docx"))
    source_document = DOCXLoader().load(str(doc_path))
    profiler = DocumentProfiler()
    profile = profiler.profile(source_document)
    inferer = TemplateInferer()
    template = inferer.infer(
        family_cluster=type("Cluster", (), {"family_id": "family_000", "document_ids": [profile.document_id], "confidence": 1.0})(),
        profiles_by_id={profile.document_id: profile},
        documents_by_id={profile.document_id: source_document},
    )

    values = {
        variable.variable_name: (variable.observed_values[0] if variable.observed_values else "sample_value")
        for variable in template.variables
    }

    output_path = tmp_path / "benchmark_regenerated.docx"
    result = DocxRegenerator().regenerate(template, values, source_document=source_document, output_path=str(output_path))

    assert result.succeeded is True
    assert output_path.exists()

    doc = DocxDocument(str(output_path))
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    assert paragraphs

    expected_titles = [section.title_or_pattern for section in template.sections if section.title_or_pattern and section.inferred_type in {"REQUIRED", "CONDITIONAL"}]
    for title in expected_titles:
        title_text = title.replace("_", " ")
        assert any(title_text in paragraph.lower() for paragraph in (p.lower() for p in paragraphs))
